#!/usr/bin/env python3
"""
Outil de remplissage des rowKey dans les pure-form XLSM
========================================================

Détecte les sheets avec Open Rows (rowNumbering=999) et remplit la cellule
rowKey en se basant sur l'analyse du DPM (Annotated Table Layout).

Patterns rowKey :
  - "Sheet per" (z-axis)  : ${sheetName|r999|cXXXX}
  - Inline key (pas z-axis): ${sheetName|r999|cXXXX};domain_dataset.datasetField

Modes CLI :
  python rowkey_tool.py                     # Rapport / analyse
  python rowkey_tool.py --write             # Écrire les rowKey
  python rowkey_tool.py --gui               # Interface graphique

  # Écrire un rowKey manuellement :
  python rowkey_tool.py --set-rowkey "formule" --sheet SHEET_NAME
"""

import argparse
import glob
import os
import re
import sys
import shutil
import openpyxl

# ── Configuration ──────────────────────────────────────────────────────────────

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SKIP_SHEETS = {'identification', 'cellSpecificities', 'labelTranslation', 'TOC'}


# ══════════════════════════════════════════════════════════════════════════════
#  UTILITAIRES
# ══════════════════════════════════════════════════════════════════════════════

def find_xlsm_files(base_dir):
    """Trouve tous les fichiers pure-form XLSM valides."""
    xlsm_files = sorted(glob.glob(os.path.join(base_dir, 'pure-form-*.xlsm')))
    return [f for f in xlsm_files
            if '~$' not in os.path.basename(f)
            and ' copy' not in os.path.basename(f).lower()]


def create_backup(fpath):
    """Crée une sauvegarde .bak si elle n'existe pas déjà."""
    backup_path = fpath + '.bak'
    if not os.path.exists(backup_path):
        shutil.copy2(fpath, backup_path)
        print(f"  Sauvegarde créée : {os.path.basename(backup_path)}")
    return backup_path


def find_header_row(ws):
    """Cherche la ligne d'en-tête (où A='level') entre les lignes 15 et 30."""
    for r in range(15, 31):
        if ws.cell(row=r, column=1).value == 'level':
            return r
    return None


def get_domain_from_identification(wb):
    """Extrait le domaine (linkedDomain) depuis la sheet identification."""
    if 'identification' not in wb.sheetnames:
        return None
    ws = wb['identification']
    # Row 1 = headers, Row 2 = values
    # Find linkedDomain column
    for c in range(1, 20):
        if ws.cell(1, c).value == 'linkedDomain':
            return ws.cell(2, c).value
    return None


# ══════════════════════════════════════════════════════════════════════════════
#  LECTURE DPM : Open Rows et <Key value>
# ══════════════════════════════════════════════════════════════════════════════

def read_dpm_open_rows(layout_path):
    """
    Lit un fichier DPM et retourne les sheets avec Open Rows et <Key value>.

    Retourne: {sheet_name: {
        'key_col_code': str,        # ex: '0005'
        'is_sheet_per': bool,       # True si B3 = "Sheet per..."
        'key_property': str,        # ex: 'qID', 'qAE', 'EN'
        'key_hierarchy': str,       # ex: 'qRAC', 'qAE', 'EN'
        'key_text': str,            # Texte complet de la cellule <Key value>
    }}
    """
    if not os.path.exists(layout_path):
        print(f"  ERREUR: Fichier DPM non trouvé: {layout_path}", file=sys.stderr)
        return {}

    wb = openpyxl.load_workbook(layout_path, read_only=True, data_only=True)
    result = {}

    for sn in wb.sheetnames:
        if sn == 'TOC':
            continue
        ws = wb[sn]

        # Check for "Open Rows" in column B
        has_open_rows = False
        for r in range(1, 25):
            b = ws.cell(r, 2).value
            if b and "Open Rows" in str(b):
                has_open_rows = True
                break

        if not has_open_rows:
            continue

        # Check if "Sheet per" pattern
        b3 = ws.cell(3, 2).value
        is_sheet_per = b3 is not None and "Sheet per" in str(b3)

        # Find ALL <Key value> occurrences
        key_cells = []  # [(row, col, text)]
        for r in range(1, 30):
            for c in range(2, 25):
                v = ws.cell(r, c).value
                if v and "<Key value>" in str(v):
                    key_cells.append((r, c, str(v)))

        if not key_cells:
            # Open Rows sans <Key value> → pas de rowKey
            result[sn] = {
                'key_col_code': None,
                'is_sheet_per': is_sheet_per,
                'key_property': None,
                'key_hierarchy': None,
                'key_text': None,
            }
            continue

        # Find the code row (row with column codes like 0005, 0010, etc.)
        code_row = None
        for r in range(5, 20):
            for c in range(4, 25):
                v = ws.cell(r, c).value
                if v and re.match(r'^0\d{3}$', str(v).strip()):
                    code_row = r
                    break
            if code_row:
                break

        # When "Sheet per" is active, <Key value> cells above the code row
        # are z-axis dimensions (not row keys). Keep only those below.
        if is_sheet_per and code_row:
            row_key_cells = [(r, c, t) for r, c, t in key_cells if r > code_row]
            if row_key_cells:
                key_cells = row_key_cells

        # Deduplicate by column
        seen_cols = set()
        unique_key_cells = []
        for r, c, t in key_cells:
            if c not in seen_cols:
                seen_cols.add(c)
                unique_key_cells.append((r, c, t))
        key_cells = unique_key_cells

        # Use first key cell
        key_row, key_col, key_text = key_cells[0]

        key_col_code = None
        if code_row:
            v = ws.cell(code_row, key_col).value
            if v:
                key_col_code = str(v).strip()

        # Extract property/hierarchy from key text
        # Patterns:
        #   (qID:) <Key value>
        #   (qAE:qAE0(qAE))  (Exposure...) <Key value>
        #   (_NA:) <Key value>
        #   (qCO:new_CO1(qCO)) <Key value>[...]
        key_property = None
        key_hierarchy = None
        if key_text:
            # Try (hierarchy:variant(hierarchy)) pattern
            m = re.match(r'\((\w+):(\w*)\(', key_text)
            if m:
                key_hierarchy = m.group(1)
                key_property = m.group(2) if m.group(2) else m.group(1)
            else:
                # Try (property:) pattern
                m = re.match(r'\((\w+):\)', key_text)
                if m:
                    key_hierarchy = m.group(1)
                    key_property = m.group(1)

        result[sn] = {
            'key_col_code': key_col_code,
            'is_sheet_per': is_sheet_per,
            'key_property': key_property,
            'key_hierarchy': key_hierarchy,
            'key_text': key_text,
        }

    wb.close()
    return result


# ══════════════════════════════════════════════════════════════════════════════
#  SCAN XLSM : sheets avec Open Rows (rowNumbering=999)
# ══════════════════════════════════════════════════════════════════════════════

def scan_xlsm_rowkeys(base_dir):
    """
    Scanne tous les pure-form XLSM et retourne les sheets avec rowNumbering=999.

    Retourne: [{
        'file': str,          # nom du fichier
        'file_path': str,     # chemin complet
        'sheet': str,         # nom de la sheet
        'header_row': int,    # ligne d'en-tête (level/rowKey/...)
        'rowkey_col': int,    # colonne du rowKey (habituellement B=2)
        'rownumbering_col': int,
        'open_rows': [(row, rowkey_value)],   # lignes avec 999
        'h2_value': str,      # allocationRuleCondition (H2)
        'domain': str,        # linkedDomain du fichier
    }]
    """
    result = []
    for fpath in find_xlsm_files(base_dir):
        try:
            wb = openpyxl.load_workbook(fpath, keep_vba=True)
        except Exception:
            continue

        domain = get_domain_from_identification(wb)
        fname = os.path.basename(fpath)

        for sn in wb.sheetnames:
            if sn in SKIP_SHEETS:
                continue
            ws = wb[sn]

            header_row = find_header_row(ws)
            if not header_row:
                continue

            # Find rowKey and rowNumbering columns
            rowkey_col = None
            rownumbering_col = None
            for c in range(1, 15):
                hv = ws.cell(header_row, c).value
                if hv:
                    hv_s = str(hv).strip()
                    if hv_s == 'rowKey':
                        rowkey_col = c
                    elif hv_s == 'rowNumbering':
                        rownumbering_col = c

            if not rownumbering_col:
                continue

            # Find rows with rowNumbering=999
            open_rows = []
            for r in range(header_row + 1, ws.max_row + 1):
                v = ws.cell(r, rownumbering_col).value
                if v is not None and str(v).strip() == '999':
                    rk_val = ws.cell(r, rowkey_col).value if rowkey_col else None
                    open_rows.append((r, rk_val))

            if not open_rows:
                continue

            # H2 = allocationRuleCondition
            h2 = ws.cell(2, 8).value

            result.append({
                'file': fname,
                'file_path': fpath,
                'sheet': sn,
                'header_row': header_row,
                'rowkey_col': rowkey_col or 2,  # default to B
                'rownumbering_col': rownumbering_col,
                'open_rows': open_rows,
                'h2_value': h2,
                'domain': domain,
            })

        wb.close()

    return result


# ══════════════════════════════════════════════════════════════════════════════
#  CONSTRUCTION ET MATCHING : DPM → XLSM → rowKey
# ══════════════════════════════════════════════════════════════════════════════

def extract_dataset_field_from_h2(h2_value):
    """
    Extrait le champ dataset.field depuis une allocationRuleCondition H2.
    Exemple: "euebacorep_prudentvaluation.euebacorepRankava <> ('NULL')"
             → "euebacorep_prudentvaluation.euebacorepRankava"
    """
    if not h2_value:
        return None
    # Pattern: domain_dataset.domainField <> ('NULL')
    m = re.search(r'(\w+\.\w+)\s*<>\s*\(\s*[\'"]NULL[\'"]\s*\)', str(h2_value))
    if m:
        return m.group(1)
    return None


def build_rowkey(sheet_name, key_col_code, dataset_field=None):
    """
    Construit la formule rowKey.
      - Simple  : ${sheet|r999|cXXXX}
      - Avec ref: ${sheet|r999|cXXXX};dataset.field
    """
    base = f"${{{sheet_name}|r999|c{key_col_code}}}"
    if dataset_field:
        return f"{base};{dataset_field}"
    return base


def match_and_prepare(xlsm_items, dpm_data):
    """
    Croise les données XLSM (sheets avec 999) et DPM (open rows / key values).
    Retourne une liste d'actions à effectuer.

    Chaque action: {
        'file_path': str,
        'file': str,
        'sheet': str,
        'rowkey_col': int,
        'rows': [(row_num, existing_value)],
        'expected': str,         # formule rowKey attendue
        'status': str,           # 'à renseigner', '✓ OK', '⚠ à corriger'
        'needs_write': bool,
        'h2_value': str,
        'h2_expected': str,      # allocationRuleCondition attendue (ou None)
        'h2_needs_write': bool,
        'key_col_code': str,
        'is_sheet_per': bool,
        'dataset_field': str,
        'dpm_source': str,       # info de diagnostic
    }
    """
    actions = []

    for item in xlsm_items:
        sheet_name = item['sheet']
        h2 = item['h2_value']

        # Match avec le DPM
        dpm_info = dpm_data.get(sheet_name)

        key_col_code = None
        is_sheet_per = False
        dataset_field = None
        dpm_source = ""
        h2_expected = None
        h2_needs_write = False

        if dpm_info:
            key_col_code = dpm_info['key_col_code']
            is_sheet_per = dpm_info['is_sheet_per']
            dpm_source = f"DPM: {dpm_info['key_text'] or 'Open Rows (no key)'}".replace('\n', ' ')[:80]

            if not is_sheet_per and key_col_code:
                # Inline key → need dataset.field from H2 or flag
                dataset_field = extract_dataset_field_from_h2(h2)
                if dataset_field:
                    h2_expected = f"{dataset_field} <> ('NULL')"
                else:
                    # H2 empty → flag as incomplete inline key
                    dataset_field = "NOT_RESOLVED"
                    print(f"  ⚠ {sheet_name}: inline key (pas 'Sheet per') mais H2 vide → dataset.field=NOT_RESOLVED")
        else:
            # Not found in DPM, try to detect from XLSM existing values
            for _, rk_val in item['open_rows']:
                if rk_val:
                    m = re.match(r'\$\{[^}]+\|r999\|c(\d+)\}', str(rk_val).strip())
                    if m:
                        key_col_code = m.group(1)
                        break
            dpm_source = "XLSM (pas de DPM)"

        if not key_col_code:
            # Cannot determine key column
            expected = None
            status = "⚠ colonne clé inconnue"
            needs_write = False
        else:
            expected = build_rowkey(sheet_name, key_col_code, dataset_field)

            # Compare with existing
            existing_values = [rk for _, rk in item['open_rows'] if rk]
            if existing_values:
                existing_clean = str(existing_values[0]).strip()
                if existing_clean == expected:
                    status = "✓ OK"
                    needs_write = False
                else:
                    status = f"⚠ à corriger ({existing_clean} → {expected})"
                    needs_write = True
            else:
                status = "à renseigner"
                needs_write = True

            # Flag incomplete inline keys
            if not is_sheet_per and dataset_field == "NOT_RESOLVED":
                status = f"⚠ inline key, H2 vide (dataset.field manquant)"
                needs_write = True

            # Check if H2 needs writing for inline key
            if not is_sheet_per and dataset_field and h2_expected:
                if h2 and h2.strip() == h2_expected.strip():
                    h2_needs_write = False
                elif not h2:
                    # H2 is empty but we know the field from... well it came FROM H2
                    # so this shouldn't happen. h2_needs_write stays False.
                    h2_needs_write = False

        actions.append({
            'file_path': item['file_path'],
            'file': item['file'],
            'sheet': sheet_name,
            'rowkey_col': item['rowkey_col'],
            'rows': item['open_rows'],
            'expected': expected,
            'status': status,
            'needs_write': needs_write,
            'h2_value': h2,
            'h2_expected': h2_expected,
            'h2_needs_write': h2_needs_write,
            'key_col_code': key_col_code,
            'is_sheet_per': is_sheet_per,
            'dataset_field': dataset_field,
            'dpm_source': dpm_source,
            'header_row': item['header_row'],
        })

    return actions


# ══════════════════════════════════════════════════════════════════════════════
#  ÉCRITURE
# ══════════════════════════════════════════════════════════════════════════════

def write_rowkey(action):
    """Écrit le rowKey dans le fichier XLSM pour une action donnée."""
    fpath = action['file_path']
    sheet_name = action['sheet']
    expected = action['expected']

    if not expected:
        print(f"  ⚠ {sheet_name}: pas de rowKey à écrire (colonne clé inconnue)")
        return False

    create_backup(fpath)

    wb = openpyxl.load_workbook(fpath, keep_vba=True)
    ws = wb[sheet_name]
    changed = False

    for row_num, existing in action['rows']:
        existing_clean = str(existing).strip() if existing else ''
        if existing_clean != expected:
            ws.cell(row_num, action['rowkey_col']).value = expected
            print(f"  ✓ {sheet_name} B{row_num}: {expected}")
            changed = True

    # Also write H2 if needed
    if action['h2_needs_write'] and action['h2_expected']:
        ws.cell(2, 8).value = action['h2_expected']
        print(f"  ✓ {sheet_name} H2: {action['h2_expected']}")
        changed = True

    if changed:
        wb.save(fpath)
        print(f"  → Sauvegardé: {os.path.basename(fpath)}")
    else:
        print(f"  — {sheet_name}: rien à modifier")

    wb.close()
    return changed


# ══════════════════════════════════════════════════════════════════════════════
#  RAPPORT (MODE CLI)
# ══════════════════════════════════════════════════════════════════════════════

def print_report(actions):
    """Affiche un rapport des rowKey détectés et attendus."""
    print("\n" + "=" * 90)
    print("  RAPPORT rowKey")
    print("=" * 90)

    to_fill = [a for a in actions if a['needs_write']]
    ok = [a for a in actions if a['status'].startswith('✓')]
    unknown = [a for a in actions if 'inconnue' in a['status']]
    incomplete = [a for a in actions if a.get('dataset_field') == 'NOT_RESOLVED']

    print(f"\n  Total sheets avec Open Rows (999): {len(actions)}")
    print(f"  ✓ OK:                              {len(ok)}")
    print(f"  À renseigner/corriger:             {len(to_fill)}")
    if unknown:
        print(f"  ⚠ Colonne clé inconnue:           {len(unknown)}")
    if incomplete:
        print(f"  ⚠ Inline key, H2 vide:            {len(incomplete)}")
        print(f"    → La formule rowKey sera écrite avec ;NOT_RESOLVED")
        print(f"    → Il faut remplir H2 (allocationRuleCondition) manuellement")

    for a in actions:
        pattern = "Sheet per" if a['is_sheet_per'] else "Inline"
        existing_display = ""
        if a['rows']:
            ex_val = a['rows'][0][1]
            if ex_val:
                existing_display = str(ex_val).strip()

        print(f"\n  [{a['status']}]")
        print(f"    Fichier : {a['file']}")
        print(f"    Sheet   : {a['sheet']}")
        print(f"    Pattern : {pattern} (col clé: {a['key_col_code'] or '?'})")
        if existing_display:
            print(f"    Actuel  : {existing_display}")
        if a['expected']:
            print(f"    Attendu : {a['expected']}")
        if a['dataset_field']:
            print(f"    Dataset : {a['dataset_field']}")
        if a['h2_value']:
            print(f"    H2      : {a['h2_value']}")
        print(f"    Source  : {a['dpm_source']}")


# ══════════════════════════════════════════════════════════════════════════════
#  MODE GUI
# ══════════════════════════════════════════════════════════════════════════════

def run_interactive(base_dir):
    """Mode interactif GUI pour analyser et écrire les rowKey."""
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.withdraw()

    # ── Étape 1 : Sélection du/des fichier(s) DPM ───────────────────────────
    messagebox.showinfo(
        "RowKey Tool",
        "Étape 1/2 : Sélectionnez le(s) fichier(s) DPM\n"
        "(Annotated Table Layout .xlsx)\n\n"
        "Vous pouvez en sélectionner plusieurs (Ctrl+clic)."
    )

    layout_paths = filedialog.askopenfilenames(
        title="Sélectionner le(s) fichier(s) DPM (Annotated Table Layout)",
        initialdir=base_dir,
        filetypes=[("Fichiers Excel", "*.xlsx"), ("Tous", "*.*")]
    )

    if not layout_paths:
        print("Annulé : Aucun fichier DPM sélectionné.")
        root.destroy()
        return

    for lp in layout_paths:
        print(f"  DPM sélectionné : {os.path.basename(lp)}")

    # ── Chargement des données ───────────────────────────────────────────────
    print("Lecture des Open Rows depuis les DPM...")
    dpm_data = {}
    for lp in layout_paths:
        partial = read_dpm_open_rows(lp)
        for sheet_name, info in partial.items():
            if sheet_name not in dpm_data:
                dpm_data[sheet_name] = info
    print(f"  {len(dpm_data)} sheets avec Open Rows trouvées dans les DPM")

    print("Scan des formulaires XLSM...")
    xlsm_items = scan_xlsm_rowkeys(base_dir)
    print(f"  {len(xlsm_items)} sheets avec rowNumbering=999 dans les XLSM")

    actions = match_and_prepare(xlsm_items, dpm_data)

    # ── Étape 2 : Fenêtre de sélection ──────────────────────────────────────
    sel_window = tk.Toplevel(root)
    sel_window.title("RowKey Tool — Remplissage des rowKey")
    sel_window.geometry("1200x600")

    title_frame = tk.Frame(sel_window, padx=10, pady=5)
    title_frame.pack(fill=tk.X)
    tk.Label(title_frame,
             text="Étape 2/2 : Sélectionnez les rowKey à écrire",
             font=("Segoe UI", 12, "bold")).pack(anchor=tk.W)
    dpm_names = ', '.join(os.path.basename(lp) for lp in layout_paths)
    tk.Label(title_frame, text=f"DPM : {dpm_names}",
             font=("Segoe UI", 9), fg="gray").pack(anchor=tk.W)

    # ── Treeview ─────────────────────────────────────────────────────────────
    tree_frame = tk.Frame(sel_window, padx=10, pady=5)
    tree_frame.pack(fill=tk.BOTH, expand=True)

    columns = ("select", "file", "sheet", "pattern", "col_code", "existing",
               "expected", "dataset_field", "status")
    tree = ttk.Treeview(tree_frame, columns=columns, show='headings', selectmode='extended')

    tree.heading("select", text="✓")
    tree.heading("file", text="Fichier")
    tree.heading("sheet", text="Sheet")
    tree.heading("pattern", text="Pattern")
    tree.heading("col_code", text="Col clé")
    tree.heading("existing", text="Actuel")
    tree.heading("expected", text="Attendu")
    tree.heading("dataset_field", text="Dataset.Field")
    tree.heading("status", text="Statut")

    tree.column("select", width=30, anchor="center")
    tree.column("file", width=220)
    tree.column("sheet", width=100)
    tree.column("pattern", width=75, anchor="center")
    tree.column("col_code", width=60, anchor="center")
    tree.column("existing", width=200)
    tree.column("expected", width=250)
    tree.column("dataset_field", width=200)
    tree.column("status", width=110)

    scroll_y = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=tree.yview)
    scroll_x = ttk.Scrollbar(tree_frame, orient=tk.HORIZONTAL, command=tree.xview)
    tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)

    tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
    scroll_x.pack(side=tk.BOTTOM, fill=tk.X)

    selected_set = set()

    def populate_tree():
        tree.delete(*tree.get_children())
        for i, a in enumerate(actions):
            pattern = "Sheet per" if a['is_sheet_per'] else "Inline"
            existing = ""
            if a['rows']:
                ex_val = a['rows'][0][1]
                if ex_val:
                    existing = str(ex_val).strip()

            sel_str = "☑" if i in selected_set else "☐"
            tag = 'selected' if i in selected_set else ''

            tree.insert("", tk.END, iid=str(i), values=(
                sel_str, a['file'], a['sheet'], pattern,
                a['key_col_code'] or '?', existing,
                a['expected'] or '—', a['dataset_field'] or '',
                a['status'],
            ), tags=(tag,))

        tree.tag_configure('selected', background='#d4edda')

    def toggle_select(event):
        item = tree.focus()
        if not item:
            return
        idx = int(item)
        if idx in selected_set:
            selected_set.discard(idx)
        else:
            if actions[idx]['expected']:
                selected_set.add(idx)
        populate_tree()
        update_count()

    # Auto-select items that need writing
    for i, a in enumerate(actions):
        if a['needs_write'] and a['expected']:
            selected_set.add(i)

    tree.bind('<Double-1>', toggle_select)
    populate_tree()

    # ── Bouton édition manuelle du dataset.field ─────────────────────────────
    edit_frame = tk.Frame(sel_window, padx=10, pady=3)
    edit_frame.pack(fill=tk.X)

    tk.Label(edit_frame, text="Dataset.Field :").pack(side=tk.LEFT)
    field_entry = tk.Entry(edit_frame, width=50)
    field_entry.pack(side=tk.LEFT, padx=5)

    def apply_field():
        item = tree.focus()
        if not item:
            messagebox.showwarning("RowKey Tool", "Sélectionnez une ligne d'abord.")
            return
        idx = int(item)
        a = actions[idx]
        field_val = field_entry.get().strip()

        if field_val:
            a['dataset_field'] = field_val
            a['expected'] = build_rowkey(a['sheet'], a['key_col_code'], field_val)
            a['h2_expected'] = f"{field_val} <> ('NULL')"
            a['h2_needs_write'] = not a['h2_value']
        else:
            a['dataset_field'] = None
            a['expected'] = build_rowkey(a['sheet'], a['key_col_code'])
            a['h2_expected'] = None
            a['h2_needs_write'] = False

        # Update status
        existing_values = [rk for _, rk in a['rows'] if rk]
        if existing_values:
            existing_clean = str(existing_values[0]).strip()
            if existing_clean == a['expected']:
                a['status'] = "✓ OK"
                a['needs_write'] = False
            else:
                a['status'] = f"⚠ à corriger"
                a['needs_write'] = True
        else:
            a['status'] = "à renseigner"
            a['needs_write'] = True

        if a['needs_write'] and a['expected']:
            selected_set.add(idx)

        populate_tree()
        update_count()

    tk.Button(edit_frame, text="Appliquer à la ligne sélectionnée",
              command=apply_field).pack(side=tk.LEFT, padx=5)

    # ── Boutons d'action ─────────────────────────────────────────────────────
    btn_frame = tk.Frame(sel_window, padx=10, pady=10)
    btn_frame.pack(fill=tk.X)

    count_label = tk.Label(btn_frame, text="", font=("Segoe UI", 9))
    count_label.pack(side=tk.LEFT, padx=10)

    def update_count():
        n = len(selected_set)
        count_label.config(text=f"{n} rowKey sélectionné(s) pour écriture")

    update_count()

    # Select All / Deselect All
    def select_all():
        for i, a in enumerate(actions):
            if a['expected']:
                selected_set.add(i)
        populate_tree()
        update_count()

    def deselect_all():
        selected_set.clear()
        populate_tree()
        update_count()

    tk.Button(btn_frame, text="Tout sélectionner", command=select_all).pack(side=tk.LEFT, padx=2)
    tk.Button(btn_frame, text="Tout désélectionner", command=deselect_all).pack(side=tk.LEFT, padx=2)

    def do_write():
        if not selected_set:
            messagebox.showwarning("RowKey Tool", "Aucun rowKey sélectionné.")
            return

        to_write = [actions[i] for i in sorted(selected_set)]
        msg = f"Écrire {len(to_write)} rowKey ?\n\n"
        for a in to_write[:10]:
            msg += f"  {a['sheet']} → {a['expected']}\n"
        if len(to_write) > 10:
            msg += f"  ... et {len(to_write) - 10} autres\n"

        if not messagebox.askyesno("Confirmer l'écriture", msg):
            return

        # Group by file
        by_file = {}
        for a in to_write:
            by_file.setdefault(a['file_path'], []).append(a)

        total_written = 0
        for fpath, file_actions in by_file.items():
            create_backup(fpath)
            wb = openpyxl.load_workbook(fpath, keep_vba=True)

            for a in file_actions:
                ws = wb[a['sheet']]
                for row_num, _ in a['rows']:
                    ws.cell(row_num, a['rowkey_col']).value = a['expected']
                    print(f"  ✓ {a['sheet']} [{row_num},{a['rowkey_col']}]: {a['expected']}")
                    total_written += 1

                if a['h2_needs_write'] and a['h2_expected']:
                    ws.cell(2, 8).value = a['h2_expected']
                    print(f"  ✓ {a['sheet']} H2: {a['h2_expected']}")

            wb.save(fpath)
            print(f"  → Sauvegardé: {os.path.basename(fpath)}")
            wb.close()

        messagebox.showinfo("RowKey Tool",
                            f"Écriture terminée !\n{total_written} rowKey écrit(s).")

        # Refresh
        nonlocal xlsm_items, actions
        xlsm_items = scan_xlsm_rowkeys(base_dir)
        actions = match_and_prepare(xlsm_items, dpm_data)
        selected_set.clear()
        populate_tree()
        update_count()

    tk.Button(btn_frame, text="✏ Écrire les rowKey sélectionnés",
              command=do_write, bg="#28a745", fg="white",
              font=("Segoe UI", 10, "bold")).pack(side=tk.RIGHT, padx=10)

    tk.Button(btn_frame, text="Annuler",
              command=lambda: (sel_window.destroy(), root.destroy())
              ).pack(side=tk.RIGHT, padx=5)

    sel_window.protocol("WM_DELETE_WINDOW",
                        lambda: (sel_window.destroy(), root.destroy()))
    root.mainloop()


# ══════════════════════════════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="Remplissage des rowKey dans les pure-form XLSM"
    )
    parser.add_argument('--gui', action='store_true',
                        help="Interface graphique")
    parser.add_argument('--write', action='store_true',
                        help="Écrire les rowKey (mode batch)")
    parser.add_argument('--set-rowkey', type=str,
                        help="Écrire manuellement un rowKey")
    parser.add_argument('--sheet', type=str,
                        help="Sheet cible (avec --set-rowkey)")
    parser.add_argument('--dpm', type=str, nargs='+',
                        help="Chemin(s) vers le(s) fichier(s) DPM")
    parser.add_argument('--dir', type=str, default=BASE_DIR,
                        help="Répertoire de travail")
    args = parser.parse_args()

    base_dir = args.dir

    if args.gui:
        run_interactive(base_dir)
        return

    if args.set_rowkey:
        # Mode écriture manuelle
        if not args.sheet:
            print("ERREUR: --sheet requis avec --set-rowkey", file=sys.stderr)
            sys.exit(1)

        # Find the sheet
        fpath = None
        for fp in find_xlsm_files(base_dir):
            try:
                wb = openpyxl.load_workbook(fp, keep_vba=True)
                if args.sheet in wb.sheetnames:
                    fpath = fp
                wb.close()
                if fpath:
                    break
            except Exception:
                continue

        if not fpath:
            print(f"ERREUR: Sheet '{args.sheet}' non trouvée", file=sys.stderr)
            sys.exit(1)

        create_backup(fpath)
        wb = openpyxl.load_workbook(fpath, keep_vba=True)
        ws = wb[args.sheet]

        header_row = find_header_row(ws)
        if not header_row:
            print(f"ERREUR: En-tête non trouvé dans {args.sheet}", file=sys.stderr)
            wb.close()
            sys.exit(1)

        # Find rowKey column
        rowkey_col = None
        rownumbering_col = None
        for c in range(1, 15):
            hv = ws.cell(header_row, c).value
            if hv:
                hv_s = str(hv).strip()
                if hv_s == 'rowKey':
                    rowkey_col = c
                elif hv_s == 'rowNumbering':
                    rownumbering_col = c

        if not rowkey_col:
            print(f"ERREUR: Colonne rowKey non trouvée dans {args.sheet}", file=sys.stderr)
            wb.close()
            sys.exit(1)

        # Find 999 rows and write
        written = 0
        for r in range(header_row + 1, ws.max_row + 1):
            if rownumbering_col:
                v = ws.cell(r, rownumbering_col).value
                if v is not None and str(v).strip() == '999':
                    ws.cell(r, rowkey_col).value = args.set_rowkey
                    print(f"  ✓ {args.sheet} [{r},{rowkey_col}]: {args.set_rowkey}")
                    written += 1

        if written:
            wb.save(fpath)
            print(f"  → Sauvegardé: {os.path.basename(fpath)} ({written} rowKey)")
        else:
            print(f"  ⚠ Aucune ligne 999 trouvée dans {args.sheet}")

        wb.close()
        return

    # Mode rapport / batch
    # Find DPM files
    dpm_paths = args.dpm
    if not dpm_paths:
        # Ouvrir un dialogue pour laisser l'utilisateur choisir le(s) DPM
        import tkinter as tk
        from tkinter import filedialog
        _root = tk.Tk()
        _root.withdraw()
        selected = filedialog.askopenfilenames(
            title="Sélectionner le(s) fichier(s) DPM (Annotated Table Layout)",
            initialdir=base_dir,
            filetypes=[("Fichiers Excel", "*.xlsx"), ("Tous", "*.*")]
        )
        _root.destroy()
        if selected:
            dpm_paths = list(selected)
        else:
            print("Annulé : Aucun fichier DPM sélectionné.", file=sys.stderr)
            sys.exit(1)

    print("Lecture des Open Rows depuis les DPM...")
    dpm_data = {}
    for lp in dpm_paths:
        partial = read_dpm_open_rows(lp)
        for sheet_name, info in partial.items():
            if sheet_name not in dpm_data:
                dpm_data[sheet_name] = info
    print(f"  {len(dpm_data)} sheets avec Open Rows trouvées dans les DPM")

    print("Scan des formulaires XLSM...")
    xlsm_items = scan_xlsm_rowkeys(base_dir)
    print(f"  {len(xlsm_items)} sheets avec rowNumbering=999 dans les XLSM")

    actions = match_and_prepare(xlsm_items, dpm_data)

    if args.write:
        to_write = [a for a in actions if a['needs_write'] and a['expected']]
        if not to_write:
            print("\n  Rien à écrire : tous les rowKey sont déjà corrects.")
            return

        print(f"\n  Écriture de {len(to_write)} rowKey...")
        by_file = {}
        for a in to_write:
            by_file.setdefault(a['file_path'], []).append(a)

        for fpath, file_actions in by_file.items():
            create_backup(fpath)
            wb = openpyxl.load_workbook(fpath, keep_vba=True)

            for a in file_actions:
                ws = wb[a['sheet']]
                for row_num, _ in a['rows']:
                    ws.cell(row_num, a['rowkey_col']).value = a['expected']
                    print(f"  ✓ {a['sheet']} [{row_num},{a['rowkey_col']}]: {a['expected']}")

                if a['h2_needs_write'] and a['h2_expected']:
                    ws.cell(2, 8).value = a['h2_expected']
                    print(f"  ✓ {a['sheet']} H2: {a['h2_expected']}")

            wb.save(fpath)
            print(f"  → Sauvegardé: {os.path.basename(fpath)}")
            wb.close()
    else:
        print_report(actions)


if __name__ == '__main__':
    main()
