#!/usr/bin/env python3
"""
Outil de récupération et renseignement des classes Z-axis
=========================================================

Ce script :
1. Lit tous les fichiers pure-form XLSM
2. Pour chaque sheet avec un Z-axis (J2 renseigné), récupère le nom du type
3. Cherche les valeurs enum correspondantes dans le dictionnaire YML
4. Mode REPORT : affiche les classes trouvées pour chaque formulaire
5. Mode WRITE  : écrit les classes dans la zone Z-axis du XLSM (lignes 5+)
6. Mode SET    : renseigne J2 (code) + J4 (associatedAttribute) + enum values
                  pour une sheet donnée

Usage :
  python zaxis_tool.py                   # Mode rapport uniquement
  python zaxis_tool.py --write           # Écrire les valeurs dans les XLSM

  # Renseigner le type Z-axis + enum pour une sheet précise :
  python zaxis_tool.py --set typspelendingexp --sheet K_29.02.a
  python zaxis_tool.py --set exposureClassSub1 --sheet K_29.00
"""

import argparse
import glob
import os
import re
import sys
import shutil
import yaml
import openpyxl

# ── Configuration ──────────────────────────────────────────────────────────────

BASE_DIR = r'c:\Users\hnasserma\Documents\test'
DICT_FILE = os.path.join(BASE_DIR, 'pure-dictionary-euebacorep.yml')
SKIP_SHEETS = {'identification', 'cellSpecificities', 'labelTranslation'}

# ── Chargement du dictionnaire ─────────────────────────────────────────────────

def load_dictionary(dict_path):
    """Charge le dictionnaire YML et construit un index des types avec leurs enum."""
    with open(dict_path, 'r', encoding='utf-8') as f:
        data = yaml.safe_load(f)

    type_index = {}  # {type_name: {enum: [...], submissionValue: {...}}}

    # Chercher dans 'enumerations' et 'definitions'
    for section_key in ('enumerations', 'definitions'):
        section = data.get(section_key, {})
        if not isinstance(section, dict):
            continue
        for type_name, entry in section.items():
            if not isinstance(entry, dict):
                continue
            enum_vals = entry.get('enum', [])
            if not enum_vals:
                continue
            sub_vals = entry.get('submissionValue', {})
            if not isinstance(sub_vals, dict):
                sub_vals = {}
            # Ne pas écraser si déjà trouvé dans une section précédente
            if type_name not in type_index:
                type_index[type_name] = {
                    'enum': enum_vals,
                    'submissionValue': sub_vals,
                    'section': section_key,
                }

    return type_index


# ── Lecture du DPM (Annotated Table Layout) ────────────────────────────────────

def read_annotated_layout(layout_path):
    """
    Lit un fichier DPM (Annotated Table Layout .xlsx) et retourne les sheets
    avec un Z-axis.
    Retourne: {sheet_name: {'hierarchy': str, 'variant': str}}
    """
    if not os.path.exists(layout_path):
        print(f"  ERREUR: Fichier DPM non trouvé: {layout_path}", file=sys.stderr)
        return {}

    wb = openpyxl.load_workbook(layout_path, read_only=True, data_only=True)
    result = {}

    for sn in wb.sheetnames:
        ws = wb[sn]
        b3 = ws.cell(row=3, column=2).value
        if not b3 or 'Sheet per' not in str(b3):
            continue

        d3 = ws.cell(row=3, column=4).value
        if not d3:
            continue

        # Parse D3: "(qAE:qAE5(qAE))  (description) <Key value>"
        m = re.search(r'\((\w+):(\w+)\(', str(d3))
        if not m:
            continue

        hierarchy = m.group(1)
        variant = m.group(2)
        result[sn] = {'hierarchy': hierarchy, 'variant': variant}

    wb.close()
    return result


# ── Mapping variant DPM → type dictionnaire ───────────────────────────────────

def map_variant_to_type(hierarchy, variant, type_index):
    """
    Mappe un code variant DPM (ex: qAE5, new_TA9, EN2) à un type dictionnaire.
    Utilise les submissionValue avec préfixe eba_<hierarchy>: pour trouver les candidats.
    Disambiguë par SubN et préfère les types où TOUTES les valeurs matchent le préfixe.
    """
    prefix = f"eba_{hierarchy}:"

    # Trouver les candidats: types dont les submissionValues utilisent ce préfixe
    candidates = []
    for type_name, info in type_index.items():
        sub_vals = info.get('submissionValue', {})
        if not sub_vals:
            continue
        has_prefix = any(str(v).startswith(prefix) for v in sub_vals.values() if v)
        if not has_prefix:
            continue
        all_prefix = all(str(v).startswith(prefix) for v in sub_vals.values() if v)
        candidates.append((type_name, all_prefix, len(info.get('enum', []))))

    if not candidates:
        return None

    # Préférer les candidats "purs" (toutes les submissionValues matchent)
    pure = [(name, count) for name, all_match, count in candidates if all_match]
    if not pure:
        pure = [(name, count) for name, _, count in candidates]

    # Essayer de trouver le nombre N dans le variant (ex: qAE5→5, qAE1→1, EN2→2)
    m = re.search(r'(\d+)$', variant.replace('new_', ''))
    if m:
        n = m.group(1)
        sub_n_pattern = f"sub{n}"
        sub_n_matches = [name for name, _ in pure if sub_n_pattern in name.lower()]
        if sub_n_matches:
            # Préférer exposureClassSub* sur les autres noms
            for name in sub_n_matches:
                if 'exposureclass' in name.lower():
                    return name
            return sub_n_matches[0]

    # Si un seul candidat pur, l'utiliser
    if len(pure) == 1:
        return pure[0][0]

    # Sinon retourner le premier candidat pur
    if pure:
        return pure[0][0]

    return None


# ── Scan des fichiers XLSM ────────────────────────────────────────────────────

def scan_xlsm_zaxis(base_dir):
    """
    Scanne tous les pure-form XLSM et retourne les infos Z-axis.
    Retourne: [{file, sheet, j2, j4, comp_rule_row}]
    """
    results = []
    xlsm_files = sorted(glob.glob(os.path.join(base_dir, 'pure-form-*.xlsm')))

    for fpath in xlsm_files:
        fname = os.path.basename(fpath)
        # Ignorer les copies
        if ' copy' in fname.lower():
            continue
        try:
            wb = openpyxl.load_workbook(fpath, keep_vba=True)
        except Exception as e:
            print(f"  ERREUR chargement {fname}: {e}", file=sys.stderr)
            continue

        for sn in wb.sheetnames:
            if sn in SKIP_SHEETS:
                continue
            ws = wb[sn]
            i1 = ws.cell(row=1, column=9).value
            if i1 != 'zAxis':
                continue

            j2 = ws.cell(row=2, column=10).value   # code (type name)
            j4 = ws.cell(row=4, column=10).value   # associatedAttribute

            # Trouver la ligne computationRuleOrder pour connaître la limite
            comp_rule_row = None
            for row in range(5, ws.max_row + 1):
                if ws.cell(row=row, column=9).value == 'computationRuleOrder':
                    comp_rule_row = row
                    break

            # Vérifier si des valeurs Z-axis existent déjà (lignes 5 à comp_rule_row-1)
            existing_values = []
            if comp_rule_row:
                for row in range(5, comp_rule_row):
                    val = ws.cell(row=row, column=10).value  # J column
                    if val:
                        existing_values.append(val)

            results.append({
                'file': fname,
                'file_path': fpath,
                'sheet': sn,
                'j2': j2,
                'j4': j4,
                'comp_rule_row': comp_rule_row,
                'existing_values': existing_values,
            })

        wb.close()

    return results


# ── Rapport ────────────────────────────────────────────────────────────────────

def print_report(zaxis_items, type_index):
    """Affiche un rapport détaillé des classes Z-axis par formulaire."""

    has_zaxis = [item for item in zaxis_items if item['j2']]
    no_zaxis = [item for item in zaxis_items if not item['j2']]

    print("\n" + "=" * 80)
    print("RAPPORT DES CLASSES Z-AXIS PAR FORMULAIRE")
    print("=" * 80)

    for item in has_zaxis:
        type_name = item['j2']
        print(f"\n┌─ Fichier : {item['file']}")
        print(f"│  Sheet   : {item['sheet']}")
        print(f"│  Type    : {type_name}")
        print(f"│  Attribut: {item['j4']}")

        if type_name in type_index:
            ti = type_index[type_name]
            enum_vals = ti['enum']
            sub_vals = ti['submissionValue']
            print(f"│  Nombre de classes : {len(enum_vals)}")
            print(f"│")
            for i, val in enumerate(enum_vals, 1):
                sv = sub_vals.get(val, '—')
                print(f"│  {i:3d}. {val:20s} → {sv}")
        else:
            print(f"│  ⚠ Type '{type_name}' NON TROUVÉ dans le dictionnaire!")

        if item['existing_values']:
            print(f"│")
            print(f"│  Valeurs déjà présentes dans le XLSM : {item['existing_values']}")

        print(f"└{'─' * 70}")

    if no_zaxis:
        print(f"\n{'─' * 80}")
        print(f"Sheets SANS Z-axis (J2 vide) : {len(no_zaxis)}")
        for item in no_zaxis:
            print(f"  - {item['file']} → [{item['sheet']}]")

    # Résumé
    print(f"\n{'=' * 80}")
    print("RÉSUMÉ")
    print(f"{'=' * 80}")
    print(f"  Total sheets avec Z-axis header : {len(zaxis_items)}")
    print(f"  Sheets avec Z-axis renseigné    : {len(has_zaxis)}")
    print(f"  Sheets sans Z-axis (normal)     : {len(no_zaxis)}")
    total_classes = sum(
        len(type_index[item['j2']]['enum'])
        for item in has_zaxis
        if item['j2'] in type_index
    )
    print(f"  Total classes à renseigner      : {total_classes}")


# ── Écriture dans les XLSM ────────────────────────────────────────────────────

def write_zaxis_values(zaxis_items, type_index, base_dir):
    """
    Écrit les valeurs enum Z-axis dans les fichiers XLSM.
    Pour chaque sheet avec Z-axis (J2 renseigné) :
      - Ligne 5 : en-tête "enumCode" / "submissionValue"
      - Lignes 6+ : une ligne par valeur enum
    """
    # Grouper par fichier pour n'ouvrir chaque XLSM qu'une fois
    by_file = {}
    for item in zaxis_items:
        if not item['j2']:
            continue
        if item['j2'] not in type_index:
            continue
        by_file.setdefault(item['file_path'], []).append(item)

    modified_count = 0

    for fpath, items in by_file.items():
        fname = os.path.basename(fpath)

        # Créer une sauvegarde
        backup_path = fpath + '.bak'
        if not os.path.exists(backup_path):
            shutil.copy2(fpath, backup_path)
            print(f"  Sauvegarde créée : {os.path.basename(backup_path)}")

        try:
            wb = openpyxl.load_workbook(fpath, keep_vba=True)
        except Exception as e:
            print(f"  ERREUR ouverture {fname}: {e}", file=sys.stderr)
            continue

        file_modified = False

        for item in items:
            sn = item['sheet']
            type_name = item['j2']
            ti = type_index[type_name]
            enum_vals = ti['enum']
            sub_vals = ti['submissionValue']

            if sn not in wb.sheetnames:
                print(f"  ⚠ Sheet {sn} non trouvée dans {fname}")
                continue

            ws = wb[sn]

            # Vérifier l'espace disponible
            comp_rule_row = item['comp_rule_row']
            start_row = 5
            needed_rows = len(enum_vals) + 1  # +1 for header row

            if comp_rule_row and (start_row + needed_rows) > comp_rule_row:
                # Pas assez d'espace - il faut décaler les lignes
                shift = (start_row + needed_rows) - comp_rule_row
                print(f"  ⚠ {fname}[{sn}] : décalage de {shift} lignes pour faire de la place")
                ws.insert_rows(comp_rule_row, amount=shift)

            # Écrire l'en-tête
            ws.cell(row=start_row, column=9, value='enumCode')
            ws.cell(row=start_row, column=10, value='submissionValue')

            # Écrire les valeurs
            for i, val in enumerate(enum_vals):
                row = start_row + 1 + i
                sv = sub_vals.get(val, '')
                ws.cell(row=row, column=9, value=val)
                ws.cell(row=row, column=10, value=sv)

            print(f"  ✓ {fname}[{sn}] : {len(enum_vals)} classes écrites (lignes {start_row+1}-{start_row+len(enum_vals)})")
            file_modified = True

        if file_modified:
            try:
                wb.save(fpath)
                modified_count += 1
                print(f"  ✓ Fichier sauvegardé : {fname}")
            except PermissionError:
                print(f"  ✗ ERREUR: '{fname}' est verrouillé (ouvert dans Excel ?). Fichier ignoré.")

        wb.close()

    print(f"\n  Total fichiers modifiés : {modified_count}")


# ── Mode SET : renseigner J2 + J4 + enum pour une sheet ───────────────────────

def set_zaxis_for_sheet(sheet_name, type_name, type_index, base_dir):
    """
    Renseigne le Z-axis complet pour une sheet donnée :
      - J2 = code (type_name)
      - J4 = associatedAttribute (euebacorep.<type_name>)
      - Lignes 5+ = enum values + submissionValue
    """
    if type_name not in type_index:
        print(f"  ERREUR: Type '{type_name}' non trouvé dans le dictionnaire!")
        print(f"  Types disponibles contenant 'expcla' ou 'exposure' :")
        for k in sorted(type_index.keys()):
            if 'expcla' in k.lower() or 'exposure' in k.lower() or 'typspe' in k.lower():
                print(f"    - {k} ({len(type_index[k]['enum'])} valeurs)")
        return False

    ti = type_index[type_name]
    enum_vals = ti['enum']
    sub_vals = ti['submissionValue']
    attribute = f"euebacorep.{type_name}"

    # Trouver le fichier XLSM contenant cette sheet
    # sheet_name = K_29.02.a → table = K_29.02
    m = re.match(r'([A-Z]_\d+\.\d+)', sheet_name)
    if not m:
        print(f"  ERREUR: Impossible d'extraire le numéro de table de '{sheet_name}'")
        return False
    table_num = m.group(1)

    xlsm_files = sorted(glob.glob(os.path.join(base_dir, 'pure-form-*.xlsm')))
    target_file = None
    for f in xlsm_files:
        fname = os.path.basename(f)
        if ' copy' in fname.lower():
            continue
        if table_num.replace('.', '_') in fname.replace('.', '_') or table_num in fname:
            # Vérifier que la sheet existe dans ce fichier
            try:
                wb_check = openpyxl.load_workbook(f, keep_vba=True)
                if sheet_name in wb_check.sheetnames:
                    target_file = f
                wb_check.close()
                if target_file:
                    break
            except Exception:
                continue

    if not target_file:
        print(f"  ERREUR: Aucun fichier XLSM trouvé contenant la sheet '{sheet_name}'")
        print(f"  Recherche pour le tableau : {table_num}")
        return False

    fname = os.path.basename(target_file)
    print(f"\n  Fichier cible : {fname}")
    print(f"  Sheet cible   : {sheet_name}")
    print(f"  Type Z-axis   : {type_name}")
    print(f"  Attribut      : {attribute}")
    print(f"  Nb classes    : {len(enum_vals)}")

    # Sauvegarde
    backup_path = target_file + '.bak'
    if not os.path.exists(backup_path):
        shutil.copy2(target_file, backup_path)
        print(f"  Sauvegarde    : {os.path.basename(backup_path)}")

    wb = openpyxl.load_workbook(target_file, keep_vba=True)
    ws = wb[sheet_name]

    # Vérifier que I1 = 'zAxis'
    i1 = ws.cell(row=1, column=9).value
    if i1 != 'zAxis':
        print(f"  ERREUR: I1 = '{i1}', attendu 'zAxis'. Cette sheet n'a pas de zone Z-axis.")
        wb.close()
        return False

    # 1. Renseigner J2 (code)
    old_j2 = ws.cell(row=2, column=10).value
    ws.cell(row=2, column=10, value=type_name)
    print(f"\n  J2 (code)               : '{old_j2}' → '{type_name}'")

    # 2. Renseigner J4 (associatedAttribute)
    old_j4 = ws.cell(row=4, column=10).value
    ws.cell(row=4, column=10, value=attribute)
    print(f"  J4 (associatedAttribute): '{old_j4}' → '{attribute}'")

    # 3. Trouver computationRuleOrder pour ne pas écraser
    comp_rule_row = None
    for row in range(5, ws.max_row + 1):
        if ws.cell(row=row, column=9).value == 'computationRuleOrder':
            comp_rule_row = row
            break
        # Aussi chercher "enumCode" d'une écriture précédente
        if ws.cell(row=row, column=9).value == 'enumCode':
            # Nettoyer ancienne écriture
            r = row
            while r < ws.max_row and ws.cell(row=r, column=9).value:
                if ws.cell(row=r, column=9).value == 'computationRuleOrder':
                    break
                ws.cell(row=r, column=9, value=None)
                ws.cell(row=r, column=10, value=None)
                r += 1

    # Re-chercher comp_rule_row après nettoyage
    comp_rule_row = None
    for row in range(5, ws.max_row + 1):
        if ws.cell(row=row, column=9).value == 'computationRuleOrder':
            comp_rule_row = row
            break

    start_row = 5
    needed_rows = len(enum_vals) + 1  # +1 pour l'en-tête

    if comp_rule_row and (start_row + needed_rows) > comp_rule_row:
        shift = (start_row + needed_rows) - comp_rule_row
        print(f"  Décalage de {shift} lignes pour faire de la place")
        ws.insert_rows(comp_rule_row, amount=shift)

    # 4. Écrire l'en-tête
    ws.cell(row=start_row, column=9, value='enumCode')
    ws.cell(row=start_row, column=10, value='submissionValue')

    # 5. Écrire les valeurs enum
    for i, val in enumerate(enum_vals):
        row = start_row + 1 + i
        sv = sub_vals.get(val, '')
        ws.cell(row=row, column=9, value=val)
        ws.cell(row=row, column=10, value=sv)

    print(f"\n  ✓ {len(enum_vals)} classes écrites (lignes {start_row+1}-{start_row+len(enum_vals)}) :")
    for i, val in enumerate(enum_vals, 1):
        sv = sub_vals.get(val, '—')
        print(f"    {i:3d}. {val:20s} → {sv}")

    wb.save(target_file)
    print(f"\n  ✓ Fichier sauvegardé : {fname}")
    wb.close()
    return True


# ── Mode interactif GUI ───────────────────────────────────────────────────────

def run_interactive(base_dir):
    """
    Mode interactif GUI :
      1. Sélection du/des fichier(s) DPM (multi-sélection)
      2. Sélection du dictionnaire
      3. Sélection des sheets (checkboxes)
      4. Écriture des J2/J4
    """
    import tkinter as tk
    from tkinter import filedialog, messagebox

    root = tk.Tk()
    root.withdraw()

    # ── Étape 1 : Sélection du/des fichier(s) DPM ───────────────────────────
    messagebox.showinfo(
        "Z-Axis Tool",
        "Étape 1/4 : Sélectionnez le(s) fichier(s) DPM\n"
        "(Annotated Table Layout .xlsx)\n\n"
        "Vous pouvez en sélectionner plusieurs (Ctrl+clic)."
    )

    layout_paths = filedialog.askopenfilenames(
        title="Sélectionner le(s) fichier(s) DPM (Annotated Table Layout)",
        initialdir=base_dir,
        filetypes=[("Fichiers Excel", "*.xlsx"), ("Tous", "*.*")]
    )

    if not layout_paths:
        print("Annulé : Aucun fichier DPM sélectionné.")
        root.destroy()
        return

    for lp in layout_paths:
        print(f"  DPM sélectionné : {os.path.basename(lp)}")

    # ── Étape 2 : Sélection du dictionnaire YML ─────────────────────────────
    messagebox.showinfo(
        "Z-Axis Tool",
        "Étape 2/4 : Sélectionnez le fichier dictionnaire\n"
        "(pure-dictionary .yml)"
    )

    dict_path = filedialog.askopenfilename(
        title="Sélectionner le fichier dictionnaire (.yml)",
        initialdir=base_dir,
        filetypes=[("Fichiers YAML", "*.yml *.yaml"), ("Tous", "*.*")]
    )

    if not dict_path:
        print("Annulé : Aucun dictionnaire sélectionné.")
        root.destroy()
        return

    print(f"Dictionnaire sélectionné : {os.path.basename(dict_path)}")

    # Charger le dictionnaire
    print("Chargement du dictionnaire...")
    type_index = load_dictionary(dict_path)
    print(f"  {len(type_index)} types trouvés")

    # Lire tous les DPM et fusionner
    print("Lecture des DPM...")
    ann_zaxis = {}
    for lp in layout_paths:
        partial = read_annotated_layout(lp)
        for sheet_name, info in partial.items():
            if sheet_name not in ann_zaxis:
                ann_zaxis[sheet_name] = info

    if ann_zaxis:
        print(f"  {len(ann_zaxis)} sheets avec Z-axis trouvées dans les DPM")
    else:
        print("  Aucune sheet avec Z-axis dans les DPM sélectionnés")

    # Construire l'index sheet → XLSM
    xlsm_files = sorted(glob.glob(os.path.join(base_dir, 'pure-form-*.xlsm')))
    xlsm_files = [f for f in xlsm_files
                  if '~$' not in os.path.basename(f)
                  and ' copy' not in os.path.basename(f).lower()]

    sheet_to_file = {}
    for fpath in xlsm_files:
        try:
            wb = openpyxl.load_workbook(fpath, keep_vba=True)
            for sn in wb.sheetnames:
                sheet_to_file[sn] = fpath
            wb.close()
        except Exception:
            continue

    # Préparer la liste des sheets sélectionnables
    selectable = []

    # Part A : Sheets trouvées dans le(s) DPM
    for ann_sheet, info in ann_zaxis.items():
        hierarchy = info['hierarchy']
        variant = info['variant']
        type_name = map_variant_to_type(hierarchy, variant, type_index)

        fpath = sheet_to_file.get(ann_sheet)

        # Lire le J2 actuel + H2 pour raffinement du type
        current_j2 = None
        h2_value = None
        j3_value = None
        if fpath:
            try:
                wb_tmp = openpyxl.load_workbook(fpath, keep_vba=True)
                if ann_sheet in wb_tmp.sheetnames:
                    ws_tmp = wb_tmp[ann_sheet]
                    current_j2 = ws_tmp.cell(row=2, column=10).value
                    h2_value = ws_tmp.cell(row=2, column=8).value
                    # Si H2 contient un type → raffiner via H2
                    if h2_value:
                        matches = re.findall(
                            r'(euebacorep\w*\.euebacorep(\w+))', str(h2_value))
                        for full_path, tok in reversed(matches):
                            for tn in type_index:
                                if tn.lower() == tok.lower():
                                    type_name = tn
                                    j3_value = full_path
                                    break
                            else:
                                continue
                            break
                wb_tmp.close()
            except Exception:
                pass

        # Déterminer le statut
        if not current_j2:
            status = "à renseigner"
            needs_write = True
        elif type_name and current_j2.lower() != type_name.lower():
            status = f"⚠ à corriger ({current_j2} → {type_name})"
            needs_write = True
        else:
            status = "✓ OK"
            needs_write = False
        selectable.append({
            'sheet': ann_sheet,
            'file': fpath,
            'type_name': type_name,
            'source': 'DPM',
            'current_j2': current_j2,
            'status': status,
            'needs_write': needs_write,
            'h2': h2_value,
            'j3': j3_value,
        })

    # Part B : Sheets XLSM avec zAxis NON présentes dans le DPM
    dpm_sheets = set(ann_zaxis.keys())
    for fpath in xlsm_files:
        try:
            wb = openpyxl.load_workbook(fpath, keep_vba=True)
            for sn in wb.sheetnames:
                if sn in SKIP_SHEETS or sn in dpm_sheets:
                    continue
                ws = wb[sn]
                if ws.cell(row=1, column=9).value != 'zAxis':
                    continue

                j2 = ws.cell(row=2, column=10).value
                j4 = ws.cell(row=4, column=10).value

                # Détecter le type depuis J2 ou J4
                type_name = None
                if j2:
                    if j2 in type_index:
                        type_name = j2
                    else:
                        j2_lower = j2.lower()
                        for tn in type_index:
                            if tn.lower() == j2_lower:
                                type_name = tn
                                break
                if not type_name and j4:
                    m = re.search(r'euebacorep\.(\w+)', str(j4))
                    if m and m.group(1) in type_index:
                        type_name = m.group(1)

                if not j2:
                    status = "à renseigner"
                    needs_write = True
                elif type_name and j2.lower() != type_name.lower():
                    status = f"⚠ à corriger ({j2} → {type_name})"
                    needs_write = True
                else:
                    status = "✓ OK"
                    needs_write = False
                selectable.append({
                    'sheet': sn,
                    'file': fpath,
                    'type_name': type_name,
                    'source': 'XLSM',
                    'current_j2': j2,
                    'status': status,
                    'needs_write': needs_write,
                })
            wb.close()
        except Exception:
            continue

    if not selectable:
        messagebox.showinfo("Z-Axis Tool", "Aucune sheet avec Z-axis trouvée.")
        root.destroy()
        return

    # ── Étape 3 : Fenêtre de sélection des sheets ───────────────────────────
    sel_window = tk.Toplevel(root)
    sel_window.title("Z-Axis Tool — Sélection des sheets")
    sel_window.geometry("950x600")

    title_frame = tk.Frame(sel_window, padx=10, pady=5)
    title_frame.pack(fill=tk.X)
    tk.Label(title_frame,
             text="Étape 3/4 : Sélectionnez les sheets à renseigner",
             font=("Segoe UI", 12, "bold")).pack(anchor=tk.W)
    dpm_names = ', '.join(os.path.basename(lp) for lp in layout_paths)
    tk.Label(title_frame, text=f"DPM : {dpm_names}",
             font=("Segoe UI", 9), fg="gray").pack(anchor=tk.W)

    # Canvas + scrollbar
    list_frame = tk.Frame(sel_window, padx=10, pady=5)
    list_frame.pack(fill=tk.BOTH, expand=True)

    canvas = tk.Canvas(list_frame)
    scrollbar = tk.Scrollbar(list_frame, orient=tk.VERTICAL, command=canvas.yview)
    inner_frame = tk.Frame(canvas)

    inner_frame.bind("<Configure>",
                     lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.create_window((0, 0), window=inner_frame, anchor=tk.NW)
    canvas.configure(yscrollcommand=scrollbar.set)

    canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    # En-tête
    headers = ["", "Sheet", "Type", "Source", "Statut", "Fichier"]
    for col, h in enumerate(headers):
        tk.Label(inner_frame, text=h, font=("Segoe UI", 9, "bold"),
                 borderwidth=1, relief="solid", padx=5, pady=2
                 ).grid(row=0, column=col, sticky="ew")

    check_vars = []
    for i, item in enumerate(selectable):
        can_select = item['file'] is not None and item['type_name'] is not None
        var = tk.BooleanVar(value=can_select and item.get('needs_write', True))
        check_vars.append(var)

        cb = tk.Checkbutton(inner_frame, variable=var,
                            state=tk.NORMAL if can_select else tk.DISABLED)
        cb.grid(row=i + 1, column=0)

        fname = os.path.basename(item['file']) if item['file'] else "—"
        type_str = item['type_name'] or "?"
        fg = "black" if can_select else "gray"

        tk.Label(inner_frame, text=item['sheet'], fg=fg, anchor="w"
                 ).grid(row=i + 1, column=1, sticky="w", padx=5)
        tk.Label(inner_frame, text=type_str, fg=fg, anchor="w"
                 ).grid(row=i + 1, column=2, sticky="w", padx=5)
        tk.Label(inner_frame, text=item['source'], fg=fg, anchor="w"
                 ).grid(row=i + 1, column=3, sticky="w", padx=5)
        tk.Label(inner_frame, text=item['status'], fg=fg, anchor="w"
                 ).grid(row=i + 1, column=4, sticky="w", padx=5)
        tk.Label(inner_frame, text=fname, fg=fg, anchor="w"
                 ).grid(row=i + 1, column=5, sticky="w", padx=5)

    # Boutons
    btn_frame = tk.Frame(sel_window, padx=10, pady=10)
    btn_frame.pack(fill=tk.X)

    selected_items = []

    def on_write():
        for i, var in enumerate(check_vars):
            if var.get():
                selected_items.append(selectable[i])
        sel_window.destroy()

    def on_cancel():
        sel_window.destroy()

    def select_all():
        for i, var in enumerate(check_vars):
            item = selectable[i]
            if item['file'] and item['type_name']:
                var.set(True)

    def deselect_all():
        for var in check_vars:
            var.set(False)

    tk.Button(btn_frame, text="Tout sélectionner",
              command=select_all).pack(side=tk.LEFT, padx=5)
    tk.Button(btn_frame, text="Tout désélectionner",
              command=deselect_all).pack(side=tk.LEFT, padx=5)
    tk.Button(btn_frame, text="Annuler",
              command=on_cancel).pack(side=tk.RIGHT, padx=5)
    tk.Button(btn_frame, text="Écrire J2/J4", command=on_write,
              bg="#4CAF50", fg="white",
              font=("Segoe UI", 10, "bold")).pack(side=tk.RIGHT, padx=5)

    sel_window.wait_window()

    if not selected_items:
        print("Aucune sheet sélectionnée.")
        root.destroy()
        return

    # ── Étape 4 : Écriture J2/J4 ────────────────────────────────────────────
    print(f"\nÉcriture de {len(selected_items)} sheets...")

    by_file = {}
    for item in selected_items:
        by_file.setdefault(item['file'], []).append(item)

    written = 0
    for fpath, items in by_file.items():
        fname = os.path.basename(fpath)

        # Backup
        backup_path = fpath + '.bak'
        if not os.path.exists(backup_path):
            shutil.copy2(fpath, backup_path)
            print(f"  Sauvegarde créée : {os.path.basename(backup_path)}")

        try:
            wb = openpyxl.load_workbook(fpath, keep_vba=True)
        except PermissionError:
            print(f"  ✗ '{fname}' verrouillé (ouvert dans Excel ?)")
            continue
        except Exception as e:
            print(f"  ✗ Erreur ouverture {fname}: {e}")
            continue

        for item in items:
            sn = item['sheet']
            type_name = item['type_name']
            if sn not in wb.sheetnames:
                print(f"  ⚠ Sheet '{sn}' non trouvée dans {fname}")
                continue

            ws = wb[sn]
            attribute = f"euebacorep.{type_name}"
            has_h2 = item.get('h2') is not None
            j3_val = item.get('j3')

            # J2 : minuscules si type vient de H2, sinon casse dictionnaire
            j2_val = type_name.lower() if has_h2 else type_name

            old_j2 = ws.cell(row=2, column=10).value
            old_j3 = ws.cell(row=3, column=10).value
            old_j4 = ws.cell(row=4, column=10).value

            ws.cell(row=2, column=10, value=j2_val)
            ws.cell(row=4, column=10, value=attribute)

            # J3 (associatedDatasetField) : renseigner si H2 fournit le chemin
            if j3_val:
                ws.cell(row=3, column=10, value=j3_val)

            log = f"  ✓ {fname}[{sn}] : J2='{old_j2}'→'{j2_val}' | J4='{old_j4}'→'{attribute}'"
            if j3_val:
                log += f" | J3='{old_j3}'→'{j3_val}'"
            print(log)
            written += 1

        try:
            wb.save(fpath)
            print(f"  ✓ Fichier sauvegardé : {fname}")
        except PermissionError:
            print(f"  ✗ '{fname}' verrouillé, impossible de sauvegarder")
        wb.close()

    print(f"\nTerminé ! {written} sheets mises à jour.")
    messagebox.showinfo("Z-Axis Tool",
                        f"Terminé !\n{written} sheets mises à jour.")
    root.destroy()


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Outil de récupération et renseignement des classes Z-axis"
    )
    parser.add_argument(
        '--write', action='store_true',
        help="Écrire les valeurs dans les fichiers XLSM (sinon : rapport seul)"
    )
    parser.add_argument(
        '--set', metavar='TYPE',
        help="Renseigner J2 (code) + J4 (attribut) + enum pour une sheet. Ex: --set typspelendingexp"
    )
    parser.add_argument(
        '--sheet', metavar='SHEET',
        help="Sheet cible pour --set. Ex: --sheet K_29.02.a"
    )
    parser.add_argument(
        '--dir', default=BASE_DIR,
        help=f"Répertoire de travail (défaut: {BASE_DIR})"
    )
    parser.add_argument(
        '--gui', action='store_true',
        help="Mode interactif GUI (sélection DPM + dictionnaire + sheets)"
    )
    args = parser.parse_args()

    base_dir = args.dir

    # ── Mode GUI ──
    if args.gui:
        run_interactive(base_dir)
        return

    dict_path = os.path.join(base_dir, 'pure-dictionary-euebacorep.yml')

    if not os.path.exists(dict_path):
        print(f"ERREUR: Dictionnaire non trouvé: {dict_path}", file=sys.stderr)
        sys.exit(1)

    # 1. Charger le dictionnaire
    print("Chargement du dictionnaire...")
    type_index = load_dictionary(dict_path)
    print(f"  {len(type_index)} types trouvés dans le dictionnaire")

    # ── Mode SET : renseigner une sheet spécifique ──
    if args.set:
        if not args.sheet:
            print("ERREUR: --set nécessite --sheet. Ex: --set typspelendingexp --sheet K_29.02.a",
                  file=sys.stderr)
            sys.exit(1)
        print(f"\n{'=' * 80}")
        print("MODE SET : Renseignement Z-axis pour une sheet")
        print(f"{'=' * 80}")
        ok = set_zaxis_for_sheet(args.sheet, args.set, type_index, base_dir)
        if ok:
            print("\nTerminé !")
        else:
            sys.exit(1)
        return

    # 2. Scanner les XLSM
    print("Scan des fichiers pure-form XLSM...")
    zaxis_items = scan_xlsm_zaxis(base_dir)
    print(f"  {len(zaxis_items)} sheets avec header zAxis trouvées")

    # 3. Rapport
    print_report(zaxis_items, type_index)

    # 4. Écriture (si demandé)
    if args.write:
        print(f"\n{'=' * 80}")
        print("ÉCRITURE DES CLASSES DANS LES FICHIERS XLSM")
        print(f"{'=' * 80}")
        write_zaxis_values(zaxis_items, type_index, base_dir)
        print("\nTerminé ! Les fichiers originaux ont été sauvegardés avec l'extension .bak")
    else:
        print(f"\n💡 Pour écrire les valeurs dans les XLSM, relancer avec : python zaxis_tool.py --write")


if __name__ == '__main__':
    main()
