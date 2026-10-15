# I need you to generate a python code to create a new Excel file named pure-dictionary-frbdfruba.xlsx 
# by reading and processing data from the source Excel file bdf_ruba_dictionnaire_v1.2.0.xlsx.
#
# Source File: bdf_ruba_dictionnaire_v1.2.0.xlsx
# Step 1 – Filter Rows in the "Metrics" Sheet
#  • Open the "Metrics" sheet.
#  • Find all rows where the "Data type" column equals: enum:enumerationItemType.
# Step 2 – For Each Matching Row
#  • Extract the "Domain" column → this is the domain name.
# Step 3 – Use the Domain Sheet
#  • For each domain, open the sheet named after the domain (matching the value from the "Domain" column).
# Step 4 – Parse Hierarchies and Enumeration Values
#  • In the domain sheet, locate the "Hierarchy" column.
#  • Identify all rows where the value starts with a number followed by a colon (e.g., 100: Description).
#  • For each such hierarchy:
#    ○ Extract the number (e.g., 100) and the description (e.g., Description).
#    ○ Find the corresponding enum values in the "Name" column (the one that is after the Hierarchy column), 
#       which appear beside each last-level element in the hierarchy.
#    ○ Compute an enum code using this formula:
#       domain_in_lowercase + hierarchy_number
#       (Example: domain GA and hierarchy 100 → ga100)
# Step 5 – Populate the New Excel File
#  • In the new Excel file pure-dictionary-frbdfruba.xlsx, fill the sheet "enumerations".
#  • For each enum code, insert a block with the following fields:
#    Field             | Value
#    ------------------|-----------------------------------------------------
#    enum code         | Computed code (e.g., ga100)
#    formatType        | "String"
#    description(EN)   | Text after the colon in the hierarchy title
#    description(FR)   | Same as above
#    value             | Enumeration values from "Name" column (one per line)
#    name(EN)          | Text in the "Hierarchy" column beside each value
#    name(FR)          | Same as above

import openpyxl
from openpyxl.styles import Border, Side
from tkinter.simpledialog import askstring
from tkinter.simpledialog import askstring
from tkinter import Tk
import os
import shutil
import os
import shutil
from tkinter.filedialog import askopenfilename, asksaveasfilename
# Prompt user for source and output file paths
def get_user_input():
    Tk().withdraw()
    # Select the empty template
    template_file = askopenfilename(
        title="Select the empty dictionary template Excel file",
        filetypes=[("Excel files", "*.xlsx *.xls")]
    )
    # Ask for regulatory code
    regulatory_code = askstring("Regulatory Code", "Enter the regulatory code (e.g., frbdfruba):")
    if not regulatory_code:
        raise ValueError("Regulatory code is required.")

    # Place the new template directly in the template/ folder
    output_file = os.path.join("template", f"pure-dictionary-{regulatory_code}.xlsx")
    # Ensure the output directory exists
    output_dir = os.path.dirname(output_file)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)
    # Copy the template file to the output file location
    shutil.copy(template_file, output_file)
    # Place the new template directly in the template/ folder
    output_file = os.path.join("template", f"pure-dictionary-{regulatory_code}.xlsx")
    # Ensure the output directory exists
    output_dir = os.path.dirname(output_file)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)
    # Copy the template file to the output file location
    shutil.copy(template_file, output_file)
    

    # Select the source Excel file
    source_file = askopenfilename(
        title="Select the source Excel file (regulatory dictionary)",
        filetypes=[("Excel files", "*.xlsx *.xls")]
    )
    return source_file, output_file

# Load both Excel workbooks
def load_workbooks(source_file, output_file):
    source_wb = openpyxl.load_workbook(source_file, data_only=True)
    output_wb = openpyxl.load_workbook(output_file)
    return source_wb, output_wb

# Identify column indices for specified headers
def get_column_indices(sheet, headers):
    indices = {}
    for col in sheet.iter_cols(min_row=1, max_row=1):
        for cell in col:
            if cell.value in headers:
                indices[cell.value] = cell.column
    return indices

# Filter rows where the data type matches the target enumeration type
def filter_rows(sheet, data_type_col, data_type_value):
    return [
        row for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[data_type_col - 1] == data_type_value
    ]

# Create a mapping of domain sheet names to attribute labels
def create_domains_dict(filtered_rows, domain_col, label_col):
    return {row[domain_col - 1]: row[label_col - 1] for row in filtered_rows}

# Collect existing enum codes to avoid duplicates
def collect_existing_enum_codes(sheet):
    existing_enum_codes = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0]:
            existing_enum_codes.add(row[0])
    return existing_enum_codes

# Process each domain sheet and extract enumeration values
def process_domain_sheets(source_wb, output_sheet, domains, existing_enum_codes):
    row_tracker = output_sheet.max_row + 1
    for domain, attribute_name in domains.items():
        domain_sheet = source_wb[domain]
        col_indices = get_column_indices(domain_sheet, ['Hierarchy', 'Name'])
        hierarchy_col = col_indices['Hierarchy']
        name_col = col_indices['Name']

        current_enum_code = current_description = None
        first_line = True
        group_start = None

        for row in domain_sheet.iter_rows(min_row=2, values_only=True):
            hierarchy = row[hierarchy_col - 1]
            name_value = row[name_col - 1]

            # Detect start of a new enumeration group
            if hierarchy and ':' in hierarchy:
                if group_start and row_tracker > group_start + 1:
                    output_sheet.row_dimensions.group(group_start + 1, row_tracker - 1, outline_level=1, hidden=True)
                number, description = hierarchy.split(':', 1)
                current_enum_code = f"{domain.lower()}{number.strip()}"
                current_description = description.strip()
                first_line = True
                group_start = row_tracker

            # Append enumeration values to the output sheet
            if current_enum_code and name_value and current_enum_code not in existing_enum_codes:
                if first_line:
                    output_sheet.append([
                        current_enum_code, 'String', current_description, current_description,
                        name_value, hierarchy, hierarchy
                    ])
                    first_line = False
                else:
                    output_sheet.append(['', '', '', '', name_value, hierarchy, hierarchy])
                row_tracker += 1

        # Group rows visually in Excel
        if group_start and row_tracker > group_start + 1:
            output_sheet.row_dimensions.group(group_start + 1, row_tracker - 1, outline_level=1, hidden=True)

# Apply thin border to all cells in the sheet
def apply_border_style(sheet):
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = thin_border

# Main function to orchestrate the workflow
def mainDictionnary():
    source_file, output_file = get_user_input()
    metrics_sheet_name = "Metrics"
    enumerations_sheet_name = "enumerations"
    source_wb, output_wb = load_workbooks(source_file, output_file)
    metrics_sheet = source_wb[metrics_sheet_name]
    enumerations_sheet = output_wb[enumerations_sheet_name]

    col_indices = get_column_indices(metrics_sheet, ['Data type', 'Domain', 'Label'])
    data_type_col = col_indices['Data type']
    domain_col = col_indices['Domain']
    label_col = col_indices['Label']

    filtered_rows = filter_rows(metrics_sheet, data_type_col, 'enum:enumerationItemType')
    domains = create_domains_dict(filtered_rows, domain_col, label_col)
    existing_enum_codes = collect_existing_enum_codes(enumerations_sheet)

    process_domain_sheets(source_wb, enumerations_sheet, domains, existing_enum_codes)
    apply_border_style(enumerations_sheet)

    output_wb.save(output_file)

if __name__ == "__main__":
    mainDictionnary() 