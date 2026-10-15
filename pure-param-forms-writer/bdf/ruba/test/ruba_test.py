import unittest
from unittest.mock import patch
import os
import sys
import openpyxl

# Ensure the main script is importable
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import dictionnary_from_dpm

def compare_excels(file1, file2, sheet_name="enumerations"):
    wb1 = openpyxl.load_workbook(file1, data_only=True)
    wb2 = openpyxl.load_workbook(file2, data_only=True)
    ws1 = wb1[sheet_name]
    ws2 = wb2[sheet_name]

    rows1 = list(ws1.iter_rows(values_only=True))
    rows2 = list(ws2.iter_rows(values_only=True))

    if len(rows1) != len(rows2):
        return False, f"Row count differs: {len(rows1)} != {len(rows2)}"

    for i, (row1, row2) in enumerate(zip(rows1, rows2), 1):
        if row1 != row2:
            return False, f"Difference at row {i}: {row1} != {row2}"

    return True, "Excels are identical"

class TestMainDictionnary(unittest.TestCase):
    @patch('dictionnary_from_dpm.askstring')
    @patch('dictionnary_from_dpm.askopenfilename')
    def test_mainDictionnary(self, mock_askopenfilename, mock_askstring):
        # Mock the file dialogs and regulatory code input
        # The order of calls: template_file, source_file
        mock_askopenfilename.side_effect = [
            os.path.abspath("template/pure-dictionary-regulatorCode.xlsx"),
            os.path.abspath("bdf/ruba/test/input/bdf_ruba_dictionnaire_v1.2.0.xlsx")
        ]
        # Mock the regulatory code input
        mock_askstring.return_value = "frbdfruba"

        # Run the main function (should not show any popup)
        dictionnary_from_dpm.mainDictionnary()

        # Check that the mocks were called as expected
        self.assertEqual(mock_askopenfilename.call_count, 2)
        mock_askstring.assert_called_once_with("Regulatory Code", "Enter the regulatory code (e.g., frbdfruba):")

        # Compare generated file with expected file
        output_file = os.path.abspath("template/pure-dictionary-frbdfruba.xlsx")
        expected_file = os.path.abspath("bdf/ruba/test/expected_result/pure-dictionary-frbdfruba.xlsx")
        result, message = compare_excels(output_file, expected_file)
        if result:
            print("Excels are identical")
        self.assertTrue(result, message)

if __name__ == "__main__":
    unittest.main()